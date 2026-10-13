"""
One-time edit: add `key`, `source_label` and `data_year` columns to the
Data_Dictionary sheet of the NourishScope workbook (plan U2, R5/R20).

- `key` is the stable, year-free field name the app uses. At each yearly
  refresh only the workbook header, `field` and `data_year` change.
- `source_label` and `data_year` build the "Source: X, YEAR" labels.

Kept in the repo as a record of how the columns were added. Running it
again is safe: it refuses to run if the columns already exist.

Usage:
  python scripts/add_dictionary_columns.py data/NourishScope_Kentucky_County_Dataset_v1.0.xlsx
"""
import shutil
import sys

import openpyxl

ACS = ("Census ACS 5-year", "2020–2024")
ERS = "USDA ERS Food Environment Atlas"
PLACES = "CDC PLACES"
ARC = "ARC"

# field -> (key, source_label, data_year)
VALUES = {
    "county_name": ("county_name", "U.S. Census Bureau", ""),
    "fips": ("fips", "U.S. Census Bureau", ""),
    "state": ("state", "U.S. Census Bureau", ""),
    "appalachian": ("appalachian", ARC, "FY2027"),
    "arc_status_fy2027": ("arc_status", ARC, "FY2027"),
    "metro_nonmetro_2023": ("metro_nonmetro", ERS, "2023"),
    "population": ("population", *ACS),
    "total_households": ("total_households", *ACS),
    "snap_households": ("snap_households", *ACS),
    "snap_household_pct": ("snap_household_pct", *ACS),
    "poverty_rate_2021": ("poverty_rate", ERS, "2021"),
    "child_poverty_rate_2021": ("child_poverty_rate", ERS, "2021"),
    "persistent_poverty_2017_21": ("persistent_poverty", ERS, "2017–2021"),
    "median_household_income_2021": ("median_household_income", ERS, "2021"),
    "diabetes_pct": ("diabetes_pct", PLACES, "2022"),
    "obesity_pct": ("obesity_pct", PLACES, "2022"),
    "high_bp_pct": ("high_bp_pct", PLACES, "2021"),
    "low_income_low_access_pct": ("low_income_low_access_pct", ERS, "2019"),
    "grocery_stores_per_1000": ("grocery_stores_per_1000", ERS, "2020"),
    "snap_retailers": ("snap_retailers", ERS, "2023"),
    "farmers_markets": ("farmers_markets", ERS, "2018"),
    "farmers_markets_snap": ("farmers_markets_snap", ERS, "2018"),
    "households_served": ("households_served", "User-entered", ""),
    "cost_per_household": ("cost_per_household", "User-entered", ""),
}

NEW_HEADERS = ["key", "source_label", "data_year"]


def main(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Data_Dictionary"]
    headers = [c.value for c in ws[1]]
    if any(h in headers for h in NEW_HEADERS):
        sys.exit("Data_Dictionary already has key/source_label/data_year columns; nothing to do.")

    field_col = headers.index("field") + 1
    start = len(headers) + 1
    for offset, name in enumerate(NEW_HEADERS):
        ws.cell(row=1, column=start + offset, value=name)

    seen = set()
    for row in range(2, ws.max_row + 1):
        field = ws.cell(row=row, column=field_col).value
        if field is None:
            continue
        if field not in VALUES:
            sys.exit(f"Unexpected field in Data_Dictionary row {row}: {field!r}")
        for offset, value in enumerate(VALUES[field]):
            ws.cell(row=row, column=start + offset, value=value or None)
        seen.add(field)

    missing = set(VALUES) - seen
    if missing:
        sys.exit(f"Data_Dictionary is missing fields: {sorted(missing)}")

    backup = path + ".bak"
    shutil.copyfile(path, backup)
    wb.save(path)
    print(f"Added {NEW_HEADERS} for {len(seen)} fields. Backup at {backup}")


if __name__ == "__main__":
    main(sys.argv[1])
